import os
import tempfile

from openpyxl import load_workbook
from rest_framework.test import APIClient

from django.test import TestCase

from karate_tests.runner import COLUMNS
from manual_test_cases.models import ManualTestCase, ManualTestStep, ManualTestSuite


class SuiteListCreateViewTests(TestCase):
    def setUp(self):
        self.client = APIClient()

    def test_create_suite_with_name(self):
        response = self.client.post('/api/manual-tests/suites/', {'name': 'Login flows'}, format='json')
        self.assertEqual(response.status_code, 201)
        body = response.json()
        self.assertEqual(body['name'], 'Login flows')
        self.assertEqual(body['test_cases'], [])

    def test_create_suite_without_name_defaults_to_blank(self):
        response = self.client.post('/api/manual-tests/suites/', {}, format='json')
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()['name'], '')

    def test_list_returns_case_and_step_counts(self):
        suite = ManualTestSuite.objects.create(name='Checkout')
        case = ManualTestCase.objects.create(suite=suite, order=1, name='Happy path')
        ManualTestStep.objects.create(test_case=case, order=1, description='Do the thing')
        ManualTestStep.objects.create(test_case=case, order=2, description='Check the thing')

        response = self.client.get('/api/manual-tests/suites/')
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(len(body), 1)
        self.assertEqual(body[0]['test_case_count'], 1)
        self.assertEqual(body[0]['step_count'], 2)


class SuiteDetailViewTests(TestCase):
    def test_returns_nested_cases_and_steps_in_order(self):
        suite = ManualTestSuite.objects.create(name='Checkout')
        case = ManualTestCase.objects.create(suite=suite, order=1, name='Happy path', environment='QA')
        ManualTestStep.objects.create(test_case=case, order=2, description='second')
        ManualTestStep.objects.create(test_case=case, order=1, description='first')

        response = APIClient().get(f'/api/manual-tests/suites/{suite.id}/')
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(len(body['test_cases']), 1)
        steps = body['test_cases'][0]['steps']
        self.assertEqual([s['description'] for s in steps], ['first', 'second'])


class CaseCreateViewTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.suite = ManualTestSuite.objects.create(name='Checkout')

    def test_missing_name_returns_400(self):
        response = self.client.post(f'/api/manual-tests/suites/{self.suite.id}/cases/', {}, format='json')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(ManualTestCase.objects.count(), 0)

    def test_blank_name_returns_400(self):
        response = self.client.post(
            f'/api/manual-tests/suites/{self.suite.id}/cases/', {'name': '   '}, format='json'
        )
        self.assertEqual(response.status_code, 400)

    def test_nonexistent_suite_returns_404(self):
        response = self.client.post('/api/manual-tests/suites/999999/cases/', {'name': 'x'}, format='json')
        self.assertEqual(response.status_code, 404)

    def test_creates_case_with_all_fields_and_auto_order(self):
        payload = {
            'name': 'Login with valid credentials',
            'description': 'Verify a valid login succeeds',
            'environment': 'QA',
            'pre_requisite': 'User exists',
            'created_by': 'Jane',
            'sprint': 'Sprint 24',
            'lob': 'Payments',
            'vertical': 'Retail',
            'feasible_for_automation': 'Yes',
            'test_case_applicability': 'Regression',
            'labels': 'smoke',
            'test_case_status': 'Active',
        }
        response = self.client.post(f'/api/manual-tests/suites/{self.suite.id}/cases/', payload, format='json')
        self.assertEqual(response.status_code, 201)
        body = response.json()
        self.assertEqual(body['order'], 1)
        for key, value in payload.items():
            self.assertEqual(body[key], value)

        second = self.client.post(
            f'/api/manual-tests/suites/{self.suite.id}/cases/', {'name': 'Second case'}, format='json'
        )
        self.assertEqual(second.json()['order'], 2)


class StepCreateViewTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.suite = ManualTestSuite.objects.create(name='Checkout')
        self.case = ManualTestCase.objects.create(suite=self.suite, order=1, name='Happy path')

    def test_all_fields_blank_returns_400(self):
        response = self.client.post(f'/api/manual-tests/cases/{self.case.id}/steps/', {}, format='json')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(ManualTestStep.objects.count(), 0)

    def test_nonexistent_case_returns_404(self):
        response = self.client.post(
            '/api/manual-tests/cases/999999/steps/', {'description': 'x'}, format='json'
        )
        self.assertEqual(response.status_code, 404)

    def test_creates_step_with_auto_order(self):
        response = self.client.post(
            f'/api/manual-tests/cases/{self.case.id}/steps/',
            {
                'description': 'Submit the login form',
                'test_data': 'user=alice, pass=secret',
                'expected_result': '200 OK, redirected to dashboard',
                'actual_result': '',
            },
            format='json',
        )
        self.assertEqual(response.status_code, 201)
        body = response.json()
        self.assertEqual(body['order'], 1)
        self.assertEqual(body['description'], 'Submit the login form')

        second = self.client.post(
            f'/api/manual-tests/cases/{self.case.id}/steps/', {'description': 'second step'}, format='json'
        )
        self.assertEqual(second.json()['order'], 2)


class ExportSuiteExcelViewTests(TestCase):
    def setUp(self):
        self.client = APIClient()

    def test_missing_excel_path_returns_400(self):
        suite = ManualTestSuite.objects.create(name='Checkout')
        ManualTestCase.objects.create(suite=suite, order=1, name='Happy path')
        response = self.client.post(f'/api/manual-tests/suites/{suite.id}/export/', {}, format='json')
        self.assertEqual(response.status_code, 400)

    def test_no_cases_returns_400(self):
        suite = ManualTestSuite.objects.create(name='Empty suite')
        response = self.client.post(
            f'/api/manual-tests/suites/{suite.id}/export/', {'excel_path': '/tmp/out.xlsx'}, format='json'
        )
        self.assertEqual(response.status_code, 400)

    def test_exports_cases_and_steps_in_karate_format(self):
        suite = ManualTestSuite.objects.create(name='Checkout')
        case1 = ManualTestCase.objects.create(
            suite=suite, order=1, name='Login with valid credentials', description='Valid login',
            environment='QA', pre_requisite='User exists', created_by='Jane', sprint='Sprint 24',
            lob='Payments', vertical='Retail', feasible_for_automation='Yes',
            test_case_applicability='Regression', labels='smoke', test_case_status='Active',
        )
        ManualTestStep.objects.create(
            test_case=case1, order=1, description='Enter valid username/password',
            test_data='user=alice, pass=secret', expected_result='Redirected to dashboard',
            actual_result='Redirected to dashboard',
        )
        ManualTestStep.objects.create(
            test_case=case1, order=2, description='Check welcome banner',
            test_data='-', expected_result='Shows "Welcome, alice"', actual_result='Shows "Welcome, alice"',
        )
        case2 = ManualTestCase.objects.create(suite=suite, order=2, name='Login with wrong password')
        # Case with no steps added yet -- still exportable as a single blank-step row.

        with tempfile.TemporaryDirectory() as d:
            excel_path = os.path.join(d, 'out.xlsx')
            response = self.client.post(
                f'/api/manual-tests/suites/{suite.id}/export/', {'excel_path': excel_path}, format='json'
            )
            self.assertEqual(response.status_code, 200)
            body = response.json()
            self.assertEqual(body['exported_case_count'], 2)
            self.assertEqual(body['exported_step_count'], 2)

            wb = load_workbook(excel_path)
            ws = wb.active
            header = [c.value for c in ws[1]]
            self.assertEqual(header, COLUMNS)

            rows = list(ws.iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(rows), 3)  # 2 steps for case 1, 1 blank row for case 2

            self.assertEqual(rows[0][0], 1)  # S.No
            self.assertEqual(rows[0][1], 'Login with valid credentials')
            self.assertEqual(rows[0][2], 'Valid login')
            self.assertEqual(rows[0][3], 'QA')
            self.assertEqual(rows[0][6], 'Enter valid username/password')
            self.assertEqual(rows[0][10], 'Jane')
            self.assertEqual(rows[0][17], 'Active')

            self.assertIsNone(rows[1][0])  # merged case-level cell, blank on the second step row
            self.assertEqual(rows[1][6], 'Check welcome banner')

            self.assertEqual(rows[2][0], 2)
            self.assertEqual(rows[2][1], 'Login with wrong password')
            self.assertIsNone(rows[2][6])  # blank step field -- openpyxl reads an empty-string cell back as None

    def test_invalid_excel_path_returns_400_without_writing(self):
        suite = ManualTestSuite.objects.create(name='Checkout')
        ManualTestCase.objects.create(suite=suite, order=1, name='Happy path')
        response = self.client.post(
            f'/api/manual-tests/suites/{suite.id}/export/', {'excel_path': 'relative/path.xlsx'}, format='json'
        )
        self.assertEqual(response.status_code, 400)
