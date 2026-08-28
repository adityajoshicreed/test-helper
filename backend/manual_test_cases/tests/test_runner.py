import os
import tempfile

from django.test import TestCase
from openpyxl import load_workbook

from karate_tests.runner import COLUMNS
from manual_test_cases.models import ManualTestCase, ManualTestStep, ManualTestSuite
from manual_test_cases.runner import export_suite_to_excel


class ExportSuiteToExcelTests(TestCase):
    def test_merges_case_level_columns_across_step_rows(self):
        suite = ManualTestSuite.objects.create(name='Suite')
        case = ManualTestCase.objects.create(suite=suite, order=1, name='Case A', environment='QA')
        ManualTestStep.objects.create(test_case=case, order=1, description='step 1')
        ManualTestStep.objects.create(test_case=case, order=2, description='step 2')

        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, 'out.xlsx')
            step_count = export_suite_to_excel([case], path)
            self.assertEqual(step_count, 2)

            wb = load_workbook(path)
            ws = wb.active
            self.assertEqual([c.value for c in ws[1]], COLUMNS)
            self.assertEqual(len(ws.merged_cells.ranges), len(_MERGED_COLUMNS_FOR_TEST))

    def test_creates_parent_directory_if_missing(self):
        suite = ManualTestSuite.objects.create(name='Suite')
        case = ManualTestCase.objects.create(suite=suite, order=1, name='Case A')
        ManualTestStep.objects.create(test_case=case, order=1, description='step 1')

        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, 'nested', 'dir', 'out.xlsx')
            export_suite_to_excel([case], path)
            self.assertTrue(os.path.isfile(path))


_MERGED_COLUMNS_FOR_TEST = (1, 2, 3, 4, 5, 11, 12, 13, 14, 15, 16, 17, 18)
