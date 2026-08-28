"""Writes hand-authored ManualTestCase/ManualTestStep rows to the same
Excel test-case format karate_tests.runner produces, so manually-authored
and Karate-report-derived test cases land in an identical sheet layout and
can be dropped into the same tracker/process downstream."""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from karate_tests.runner import COLUMNS

# Columns that are constant for a whole test case (not per step) get merged
# into a single spanning cell across that case's step rows -- same set as
# karate_tests.runner.write_excel: S.No, Test Case Name, Test Case
# Description, Environment, Pre-Requisite, Created By, Sprint, LOB,
# Vertical, Feasible for Automation?, Test Case Applicability, Labels, Status.
_MERGED_COLUMNS = (1, 2, 3, 4, 5, 11, 12, 13, 14, 15, 16, 17, 18)
_COLUMN_WIDTHS = [8, 30, 30, 14, 20, 8, 20, 60, 45, 45, 14, 12, 12, 14, 18, 20, 16, 12]


def export_suite_to_excel(cases, excel_path):
    """`cases` is an iterable of ManualTestCase rows (its `.steps` related
    manager is queried directly). Returns the number of step rows written."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Test Cases'
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    wrap = Alignment(wrap_text=True, vertical='top')
    merged_alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')

    step_count = 0
    for case_number, case in enumerate(cases, start=1):
        start_row = ws.max_row + 1
        steps = list(case.steps.all()) or [None]
        for step_number, step in enumerate(steps, start=1):
            is_first_step = step_number == 1
            ws.append([
                case_number if is_first_step else None,
                case.name if is_first_step else None,
                case.description if is_first_step else None,
                case.environment if is_first_step else None,
                case.pre_requisite if is_first_step else None,
                step_number,
                step.description if step else '',
                step.test_data if step else '',
                step.expected_result if step else '',
                step.actual_result if step else '',
                case.created_by if is_first_step else None,
                case.sprint if is_first_step else None,
                case.lob if is_first_step else None,
                case.vertical if is_first_step else None,
                case.feasible_for_automation if is_first_step else None,
                case.test_case_applicability if is_first_step else None,
                case.labels if is_first_step else None,
                case.test_case_status if is_first_step else None,
            ])
            for cell in ws[ws.max_row]:
                cell.alignment = wrap
            if step is not None:
                step_count += 1

        end_row = ws.max_row
        if end_row > start_row:
            for col in _MERGED_COLUMNS:
                ws.merge_cells(start_row=start_row, start_column=col, end_row=end_row, end_column=col)
                ws.cell(row=start_row, column=col).alignment = merged_alignment

    for i, width in enumerate(_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

    parent = os.path.dirname(excel_path)
    if parent:
        os.makedirs(parent, exist_ok=True)
    wb.save(excel_path)
    return step_count
