import json
import os
import shutil
import tempfile
from unittest.mock import Mock, patch

import requests
from django.test import SimpleTestCase, TestCase
from openpyxl import load_workbook

from karate_ai_tests import runner
from karate_ai_tests.models import KarateAiTestCaseJob

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), 'fixtures')
SAMPLE_REPORT = os.path.join(FIXTURES_DIR, 'simple_feature_report.html')
OVERVIEW_REPORT = os.path.join(FIXTURES_DIR, 'overview_no_data.html')


def ollama_response(payload):
    """Builds a Mock standing in for requests.post()'s return value, with
    .json() -> {'response': json.dumps(payload)} the way Ollama's
    /api/generate replies when format='json' is honored."""
    mock = Mock()
    mock.raise_for_status = Mock()
    mock.json.return_value = {'response': json.dumps(payload)}
    return mock


class ValidateOllamaConfigTests(SimpleTestCase):
    def test_base_url_defaults_when_blank(self):
        self.assertEqual(runner.validate_ollama_base_url(''), runner.DEFAULT_OLLAMA_BASE_URL)

    def test_base_url_strips_trailing_slash(self):
        self.assertEqual(runner.validate_ollama_base_url('http://localhost:11434/'), 'http://localhost:11434')

    def test_base_url_rejects_non_http_scheme(self):
        with self.assertRaises(runner.PreflightError):
            runner.validate_ollama_base_url('ftp://localhost:11434')

    def test_model_rejects_empty(self):
        with self.assertRaises(runner.PreflightError):
            runner.validate_ollama_model('  ')

    def test_model_accepts_and_strips(self):
        self.assertEqual(runner.validate_ollama_model(' llama3.1 '), 'llama3.1')


class EnhanceScenarioWithAiTests(SimpleTestCase):
    def _steps(self):
        return [
            {'method': 'POST', 'url': 'http://x/api/todos', 'status_code': '201', 'response_body': '{"id": 1}'},
            {'method': 'GET', 'url': 'http://x/api/todos/1', 'status_code': '200', 'response_body': '{"id": 1}'},
        ]

    @patch('karate_ai_tests.runner.requests.post')
    def test_parses_well_formed_response(self, mock_post):
        mock_post.return_value = ollama_response({
            'steps': [
                {'step': 1, 'description': 'Create a new todo item.', 'expected_result': 'Returns 201 with the new id.'},
                {'step': 2, 'description': 'Fetch the todo just created.', 'expected_result': 'Returns 200 with matching id.'},
            ]
        })
        results = runner.enhance_scenario_with_ai('http://localhost:11434', 'llama3.1', 'a scenario', self._steps())
        self.assertEqual(results[0]['description'], 'Create a new todo item.')
        self.assertEqual(results[1]['expected_result'], 'Returns 200 with matching id.')

    @patch('karate_ai_tests.runner.requests.post')
    def test_strips_markdown_code_fence_before_parsing(self, mock_post):
        payload = {'steps': [
            {'step': 1, 'description': 'a', 'expected_result': 'b'},
            {'step': 2, 'description': 'c', 'expected_result': 'd'},
        ]}
        mock = Mock()
        mock.raise_for_status = Mock()
        mock.json.return_value = {'response': f'```json\n{json.dumps(payload)}\n```'}
        mock_post.return_value = mock
        results = runner.enhance_scenario_with_ai('http://localhost:11434', 'llama3.1', 'a scenario', self._steps())
        self.assertEqual(results[0]['description'], 'a')

    @patch('karate_ai_tests.runner.requests.post')
    def test_connection_error_raises_ai_enhancement_error(self, mock_post):
        mock_post.side_effect = requests.ConnectionError('refused')
        with self.assertRaises(runner.AiEnhancementError):
            runner.enhance_scenario_with_ai('http://localhost:11434', 'llama3.1', 'a scenario', self._steps())

    @patch('karate_ai_tests.runner.requests.post')
    def test_timeout_raises_ai_enhancement_error(self, mock_post):
        mock_post.side_effect = requests.Timeout('too slow')
        with self.assertRaises(runner.AiEnhancementError):
            runner.enhance_scenario_with_ai('http://localhost:11434', 'llama3.1', 'a scenario', self._steps())

    @patch('karate_ai_tests.runner.requests.post')
    def test_malformed_json_raises_ai_enhancement_error(self, mock_post):
        mock = Mock()
        mock.raise_for_status = Mock()
        mock.json.return_value = {'response': 'not json at all'}
        mock_post.return_value = mock
        with self.assertRaises(runner.AiEnhancementError):
            runner.enhance_scenario_with_ai('http://localhost:11434', 'llama3.1', 'a scenario', self._steps())

    @patch('karate_ai_tests.runner.requests.post')
    def test_wrong_step_count_raises_ai_enhancement_error(self, mock_post):
        mock_post.return_value = ollama_response({
            'steps': [{'step': 1, 'description': 'only one', 'expected_result': 'x'}]
        })
        with self.assertRaises(runner.AiEnhancementError):
            runner.enhance_scenario_with_ai('http://localhost:11434', 'llama3.1', 'a scenario', self._steps())

    @patch('karate_ai_tests.runner.requests.post')
    def test_missing_steps_key_raises_ai_enhancement_error(self, mock_post):
        mock_post.return_value = ollama_response({'not_steps': []})
        with self.assertRaises(runner.AiEnhancementError):
            runner.enhance_scenario_with_ai('http://localhost:11434', 'llama3.1', 'a scenario', self._steps())


class EnhanceTestCasesWithAiTests(SimpleTestCase):
    def _cases(self):
        return [
            {'name': 'case A', 'steps': [
                {'method': 'GET', 'url': 'http://x/a', 'status_code': '200', 'response_body': '{}'},
            ]},
            {'name': 'case B', 'steps': [
                {'method': 'GET', 'url': 'http://x/b', 'status_code': '200', 'response_body': '{}'},
            ]},
        ]

    @patch('karate_ai_tests.runner.enhance_scenario_with_ai')
    def test_successful_scenarios_get_ai_fields_and_are_counted(self, mock_enhance):
        mock_enhance.return_value = [{'description': 'desc', 'expected_result': 'ok'}]
        cases = self._cases()
        warnings = []
        count = runner.enhance_test_cases_with_ai(cases, 'http://localhost:11434', 'llama3.1', warnings=warnings)
        self.assertEqual(count, 2)
        self.assertEqual(cases[0]['steps'][0]['ai_description'], 'desc')
        self.assertEqual(warnings, [])

    @patch('karate_ai_tests.runner.enhance_scenario_with_ai')
    def test_one_scenario_failing_falls_back_without_affecting_the_other(self, mock_enhance):
        def side_effect(base_url, model, name, steps):
            if name == 'case A':
                raise runner.AiEnhancementError('boom')
            return [{'description': 'desc', 'expected_result': 'ok'}]

        mock_enhance.side_effect = side_effect
        cases = self._cases()
        warnings = []
        count = runner.enhance_test_cases_with_ai(cases, 'http://localhost:11434', 'llama3.1', warnings=warnings)
        self.assertEqual(count, 1)
        self.assertIsNone(cases[0]['steps'][0]['ai_description'])
        self.assertEqual(cases[1]['steps'][0]['ai_description'], 'desc')
        self.assertEqual(len(warnings), 1)
        self.assertIn('case A', warnings[0])
        self.assertIn('boom', warnings[0])


class GenerateTests(TestCase):
    def _make_job(self, reports_dir, excel_path, **overrides):
        defaults = dict(
            reports_dir=reports_dir,
            excel_path=excel_path,
            ollama_base_url='http://localhost:11434',
            ollama_model='llama3.1',
            environment='QA',
            pre_requisite='User is logged in',
            created_by='Ada',
            sprint='Sprint 1',
            lob='Payments',
            vertical='Retail',
            feasible_for_automation='Yes',
            test_case_applicability='Regression',
            labels='smoke',
            test_case_status='Active',
            status=KarateAiTestCaseJob.STATUS_RUNNING,
        )
        defaults.update(overrides)
        return KarateAiTestCaseJob.objects.create(**defaults)

    @patch('karate_ai_tests.runner.requests.post')
    def test_end_to_end_with_successful_ai_enhancement(self, mock_post):
        mock_post.return_value = ollama_response({
            'steps': [{'step': i, 'description': f'meaningful step {i}', 'expected_result': f'meaningful result {i}'}
                      for i in range(1, 6)]
        })
        with tempfile.TemporaryDirectory() as reports_dir, tempfile.TemporaryDirectory() as out_dir:
            shutil.copy(SAMPLE_REPORT, os.path.join(reports_dir, 'simple.html'))
            excel_path = os.path.join(out_dir, 'cases.xlsx')
            job = self._make_job(reports_dir, excel_path)

            runner.generate(job)
            job.refresh_from_db()

            self.assertEqual(job.status, KarateAiTestCaseJob.STATUS_COMPLETED)
            self.assertEqual(job.step_count, 5)
            self.assertEqual(job.ai_enhanced_scenario_count, 1)
            self.assertEqual(job.warnings, [])

            wb = load_workbook(excel_path)
            ws = wb.active
            header = [c.value for c in ws[1]]
            self.assertEqual(header, runner.COLUMNS)

            rows = list(ws.iter_rows(min_row=2, values_only=True))
            self.assertEqual(rows[0][6], 'meaningful step 1')  # Step Description
            self.assertIn('curl -X POST', rows[0][7])  # Test Data unchanged
            self.assertEqual(rows[0][8], 'meaningful result 1')  # Expected Result
            self.assertIn('Response Code: 201', rows[0][9])  # Actual Result still raw

    @patch('karate_ai_tests.runner.requests.post')
    def test_end_to_end_falls_back_when_ollama_is_unreachable(self, mock_post):
        mock_post.side_effect = requests.ConnectionError('refused')
        with tempfile.TemporaryDirectory() as reports_dir, tempfile.TemporaryDirectory() as out_dir:
            shutil.copy(SAMPLE_REPORT, os.path.join(reports_dir, 'simple.html'))
            excel_path = os.path.join(out_dir, 'cases.xlsx')
            job = self._make_job(reports_dir, excel_path)

            runner.generate(job)
            job.refresh_from_db()

            # The job still completes -- AI failure degrades output, it
            # doesn't fail the job.
            self.assertEqual(job.status, KarateAiTestCaseJob.STATUS_COMPLETED)
            self.assertEqual(job.ai_enhanced_scenario_count, 0)
            self.assertEqual(len(job.warnings), 1)
            self.assertIn('AI enhancement failed', job.warnings[0])

            wb = load_workbook(excel_path)
            ws = wb.active
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            self.assertEqual(rows[0][6], 'Execute the CURL')  # default fallback text
            self.assertIn('Response Code: 201', rows[0][8])  # expected falls back to raw actual result

    def test_no_html_files_marks_failed_without_calling_ai(self):
        with tempfile.TemporaryDirectory() as reports_dir, tempfile.TemporaryDirectory() as out_dir:
            excel_path = os.path.join(out_dir, 'cases.xlsx')
            job = self._make_job(reports_dir, excel_path)
            runner.generate(job)
            job.refresh_from_db()
            self.assertEqual(job.status, KarateAiTestCaseJob.STATUS_FAILED)
            self.assertIn('No Karate feature HTML reports found', job.error)
